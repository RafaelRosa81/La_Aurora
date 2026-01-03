import argparse
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from pandas.tseries.offsets import MonthEnd

VERSION = "1.0.0"
TIMESTAMP_CANDIDATES = ["timestamp", "ts", "datetime", "fechahora"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validación de integridad temporal para datos CSV."
    )
    parser.add_argument("--input-dir", required=True, help="Carpeta base con CSVs")
    parser.add_argument("--group", required=True, help="Etiqueta de grupo (estanques|bombas)")
    parser.add_argument("--asset", help="Filtra por asset (nombre o parte)")
    parser.add_argument(
        "--all",
        action="store_true",
        default=False,
        help="Procesa todos los assets detectados (default)",
    )
    parser.add_argument("--start-date", help="Fecha de inicio (YYYY-MM-DD)")
    parser.add_argument("--end-date", help="Fecha de fin (YYYY-MM-DD)")
    parser.add_argument("--freq-minutes", type=int, default=1, help="Frecuencia esperada")
    parser.add_argument(
        "--output", help="Ruta de salida para el Excel (default reports/...)"
    )
    return parser.parse_args()


def parse_date(date_text: str, is_end: bool, freq_minutes: int) -> pd.Timestamp | None:
    if not date_text:
        return None
    parsed = pd.to_datetime(date_text, errors="coerce")
    if pd.isna(parsed):
        return None
    if is_end and len(date_text) == 10:
        parsed = parsed + pd.Timedelta(days=1) - pd.Timedelta(minutes=freq_minutes)
    return parsed


def find_timestamp_column(df: pd.DataFrame) -> str | None:
    lower_map = {col.lower(): col for col in df.columns}
    for candidate in TIMESTAMP_CANDIDATES:
        if candidate in lower_map:
            return lower_map[candidate]
    return None


def determine_asset_series(df: pd.DataFrame, file_path: Path) -> pd.Series:
    lower_map = {col.lower(): col for col in df.columns}
    fallback = file_path.parent.name or file_path.stem.split("_")[0]
    asset_col = lower_map.get("asset_label")
    if asset_col:
        series = df[asset_col]
        if series.notna().any():
            series = series.astype(str).str.strip()
            series = series.where(series != "", fallback)
            return series
    return pd.Series([fallback] * len(df), index=df.index)


def load_csv(file_path: Path, freq_minutes: int) -> pd.DataFrame | None:
    try:
        df = pd.read_csv(file_path)
    except Exception as exc:
        print(f"[warning] No se pudo leer {file_path}: {exc}")
        return None

    timestamp_col = find_timestamp_column(df)
    if not timestamp_col:
        print(f"[warning] Sin columna timestamp en {file_path}")
        return None

    df["timestamp"] = pd.to_datetime(df[timestamp_col], errors="coerce")
    invalid_count = int(df["timestamp"].isna().sum())
    if invalid_count:
        print(
            f"[warning] {file_path} tiene {invalid_count} filas con timestamp inválido"
        )
    df = df[df["timestamp"].notna()].copy()
    if df.empty:
        return None

    df["asset_id"] = determine_asset_series(df, file_path)
    return df[["asset_id", "timestamp"]]


def expected_count(
    start: pd.Timestamp | None, end: pd.Timestamp | None, freq_minutes: int
) -> int:
    if not start or not end or end < start:
        return 0
    freq = pd.Timedelta(minutes=freq_minutes)
    return int((end - start) // freq) + 1


def compute_gaps(
    timestamps: pd.Series, freq_minutes: int
) -> tuple[list[dict], float]:
    freq = pd.Timedelta(minutes=freq_minutes)
    ordered = timestamps.reset_index(drop=True)
    diffs = ordered.diff()
    gaps = []
    max_gap_minutes = 0.0
    for idx, diff in diffs.items():
        if pd.isna(diff) or diff <= freq:
            continue
        gap_start = ordered.iloc[idx - 1]
        gap_end = ordered.iloc[idx]
        gap_minutes = diff.total_seconds() / 60.0
        missing_points = max(int(diff // freq) - 1, 0)
        gaps.append(
            {
                "gap_start": gap_start,
                "gap_end": gap_end,
                "gap_minutes": gap_minutes,
                "missing_points_est": missing_points,
            }
        )
        max_gap_minutes = max(max_gap_minutes, gap_minutes)
    gaps_sorted = sorted(gaps, key=lambda item: item["gap_minutes"], reverse=True)
    return gaps_sorted[:10], max_gap_minutes


def compute_duplicates(timestamps: pd.Series) -> tuple[pd.DataFrame, int]:
    counts = timestamps.value_counts()
    dup_counts = counts[counts > 1]
    duplicates_df = (
        dup_counts.rename_axis("timestamp").reset_index(name="count").sort_values("timestamp")
    )
    dup_total = int((dup_counts - 1).sum())
    return duplicates_df, dup_total


def monthly_summary(
    df: pd.DataFrame,
    range_start: pd.Timestamp | None,
    range_end: pd.Timestamp | None,
    freq_minutes: int,
) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["month", "obs", "expected", "missing_pct"])
    data = df.copy()
    data["month"] = data["timestamp"].dt.to_period("M").astype(str)
    months = sorted(data["month"].unique())
    rows = []
    for month in months:
        month_start = pd.Timestamp(f"{month}-01")
        month_end = month_start + MonthEnd(0)
        range_month_start = max(month_start, range_start) if range_start else month_start
        range_month_end = min(month_end, range_end) if range_end else month_end
        if range_month_end < range_month_start:
            continue
        observed = int(
            data[
                (data["timestamp"] >= month_start) & (data["timestamp"] <= month_end)
            ]["timestamp"].count()
        )
        expected = expected_count(range_month_start, range_month_end, freq_minutes)
        missing = max(expected - observed, 0)
        missing_pct = (missing / expected * 100.0) if expected else 0.0
        rows.append(
            {
                "month": month,
                "obs": observed,
                "expected": expected,
                "missing_pct": missing_pct,
            }
        )
    return pd.DataFrame(rows)


def build_report(
    combined: pd.DataFrame,
    group: str,
    asset_filter: str | None,
    start_date: pd.Timestamp | None,
    end_date: pd.Timestamp | None,
    freq_minutes: int,
    output_path: Path,
    input_dir: Path,
) -> None:
    summary_rows = []
    gaps_rows = []
    duplicates_rows = []
    monthly_rows = []

    if asset_filter:
        combined = combined[
            combined["asset_id"].str.contains(asset_filter, case=False, na=False)
        ]

    for asset_id, asset_df in combined.groupby("asset_id"):
        asset_df = asset_df.sort_values("timestamp")
        if start_date:
            asset_df = asset_df[asset_df["timestamp"] >= start_date]
        if end_date:
            asset_df = asset_df[asset_df["timestamp"] <= end_date]
        if asset_df.empty:
            print(f"[warning] Asset {asset_id} sin datos en el rango")
            continue

        observed = int(asset_df["timestamp"].count())
        date_min = asset_df["timestamp"].min()
        date_max = asset_df["timestamp"].max()

        range_start = start_date or date_min
        range_end = end_date or date_max
        expected = expected_count(range_start, range_end, freq_minutes)
        missing = max(expected - observed, 0)
        missing_pct = (missing / expected * 100.0) if expected else 0.0

        duplicates_df, dup_total = compute_duplicates(asset_df["timestamp"])
        if dup_total:
            examples = duplicates_df.head(5)["timestamp"].dt.strftime("%Y-%m-%d %H:%M:%S").tolist()
            print(f"[info] Asset {asset_id} duplicados={dup_total}, ejemplos={examples}")

        gaps, max_gap_minutes = compute_gaps(asset_df["timestamp"], freq_minutes)

        summary_rows.append(
            {
                "asset": asset_id,
                "obs": observed,
                "expected": expected,
                "missing_pct": missing_pct,
                "dup_count": dup_total,
                "max_gap_minutes": max_gap_minutes,
                "date_min": date_min,
                "date_max": date_max,
            }
        )

        for gap in gaps:
            gaps_rows.append(
                {
                    "asset": asset_id,
                    **gap,
                }
            )

        for _, row in duplicates_df.iterrows():
            duplicates_rows.append(
                {
                    "asset": asset_id,
                    "timestamp": row["timestamp"],
                    "count": int(row["count"]),
                }
            )

        monthly_df = monthly_summary(asset_df, range_start, range_end, freq_minutes)
        if not monthly_df.empty:
            monthly_df.insert(0, "asset", asset_id)
            monthly_rows.extend(monthly_df.to_dict("records"))

    summary_df = pd.DataFrame(
        summary_rows,
        columns=[
            "asset",
            "obs",
            "expected",
            "missing_pct",
            "dup_count",
            "max_gap_minutes",
            "date_min",
            "date_max",
        ],
    )
    gaps_df = pd.DataFrame(
        gaps_rows, columns=["asset", "gap_start", "gap_end", "gap_minutes", "missing_points_est"]
    )
    duplicates_df = pd.DataFrame(
        duplicates_rows, columns=["asset", "timestamp", "count"]
    )
    monthly_df = pd.DataFrame(
        monthly_rows, columns=["asset", "month", "obs", "expected", "missing_pct"]
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        gaps_df.to_excel(writer, sheet_name="Gaps", index=False)
        duplicates_df.to_excel(writer, sheet_name="Duplicates", index=False)
        monthly_df.to_excel(writer, sheet_name="Monthly", index=False)

        notes = [
            ("group", group),
            ("input_dir", str(input_dir)),
            ("asset_filter", asset_filter or "ALL"),
            ("start_date", start_date),
            ("end_date", end_date),
            ("freq_minutes", freq_minutes),
            ("output", str(output_path)),
            ("version", VERSION),
            ("generated_at", datetime.now()),
        ]
        notes_df = pd.DataFrame(notes, columns=["key", "value"])
        notes_df.to_excel(writer, sheet_name="Notes", index=False)

        workbook = writer.book
        if not summary_df.empty:
            ws_summary = writer.sheets["Summary"]
            missing_col = summary_df.columns.get_loc("missing_pct") + 1
            data_ref = Reference(
                ws_summary,
                min_col=missing_col,
                min_row=1,
                max_row=len(summary_df) + 1,
            )
            categories_ref = Reference(
                ws_summary,
                min_col=1,
                min_row=2,
                max_row=len(summary_df) + 1,
            )
            chart = BarChart()
            chart.title = "Missing % por asset"
            chart.y_axis.title = "% Missing"
            chart.x_axis.title = "Asset"
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(categories_ref)
            ws_summary.add_chart(chart, "J2")

        if not monthly_df.empty and monthly_df["month"].nunique() > 1:
            ws_monthly = writer.sheets["Monthly"]
            pivot = (
                monthly_df.pivot(index="month", columns="asset", values="missing_pct")
                .reset_index()
                .fillna(0.0)
            )
            start_row = len(monthly_df) + 3
            pivot.to_excel(writer, sheet_name="Monthly", index=False, startrow=start_row)
            ws_monthly = writer.sheets["Monthly"]
            max_row = start_row + len(pivot)
            max_col = len(pivot.columns)
            data_ref = Reference(
                ws_monthly,
                min_col=2,
                min_row=start_row + 1,
                max_col=max_col,
                max_row=max_row,
            )
            categories_ref = Reference(
                ws_monthly,
                min_col=1,
                min_row=start_row + 2,
                max_row=max_row,
            )
            chart = LineChart()
            chart.title = "Missing % mensual"
            chart.y_axis.title = "% Missing"
            chart.x_axis.title = "Mes"
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(categories_ref)
            ws_monthly.add_chart(chart, "J2")

        for sheet_name, df in {
            "Summary": summary_df,
            "Gaps": gaps_df,
            "Duplicates": duplicates_df,
            "Monthly": monthly_df,
            "Notes": notes_df,
        }.items():
            ws = writer.sheets[sheet_name]
            for idx, column in enumerate(df.columns, start=1):
                max_len = max(
                    [len(str(column))]
                    + [len(str(value)) for value in df[column].head(50).tolist()]
                )
                ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 40)

    print(f"[info] Reporte generado en {output_path}")


def main() -> None:
    args = parse_args()
    input_dir = Path(args.input_dir)
    if not input_dir.exists():
        raise SystemExit(f"Input dir no existe: {input_dir}")

    if args.asset and args.all:
        raise SystemExit("No uses --asset junto con --all")

    freq_minutes = max(args.freq_minutes, 1)
    start_date = parse_date(args.start_date, is_end=False, freq_minutes=freq_minutes)
    end_date = parse_date(args.end_date, is_end=True, freq_minutes=freq_minutes)

    csv_files = sorted(input_dir.rglob("*.csv"))
    if not csv_files:
        print(f"[warning] No se encontraron CSVs en {input_dir}")

    combined_frames = []
    for csv_file in csv_files:
        df = load_csv(csv_file, freq_minutes)
        if df is None or df.empty:
            continue
        combined_frames.append(df)

    if combined_frames:
        combined = pd.concat(combined_frames, ignore_index=True)
    else:
        combined = pd.DataFrame(columns=["asset_id", "timestamp"])

    if args.output:
        output_path = Path(args.output)
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = Path("reports") / f"integrity_{args.group}_{timestamp}.xlsx"

    build_report(
        combined=combined,
        group=args.group,
        asset_filter=args.asset,
        start_date=start_date,
        end_date=end_date,
        freq_minutes=freq_minutes,
        output_path=output_path,
        input_dir=input_dir,
    )


if __name__ == "__main__":
    main()
